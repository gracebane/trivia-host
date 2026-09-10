#!/usr/bin/env python3
"""birthday_collate.py — Reilly Birthday Mode: merge everyone's two-slide trivia
decks into one presentation, with an index slide the app can read.

    python3 scripts/birthday_collate.py ~/Desktop/reilly_questions
    python3 scripts/birthday_collate.py decks/ --order alpha
    python3 scripts/birthday_collate.py decks/ --order "Bob,Alice,Carol" --last Reilly
    python3 scripts/birthday_collate.py decks/ --seed 42 --points 10 --title "Reilly turns 30"

Each guest brings a .pptx named after themselves ("Alice.pptx", "bob - trivia.pptx",
"Carol's question.pptx") holding a question slide and an answer slide. This
script:

  1. takes the participant's name from the filename
  2. decides the order (random by default, seeded and printed, so it's fair AND
     reproducible; or alphabetical, file order, or an explicit list)
  3. builds ONE deck: slide 1 is an index of who's up when, then each guest's
     question slide followed by their answer slide
  4. writes pack.json alongside, in the host app's format, with NO question or
     answer text — just author, slide numbers and points. The host is playing
     too and must not be able to read the questions in the app.

The index slide carries a small machine-readable footer, so loading the merged
.pptx straight into the host app also works: the app sees the marker, switches
to Birthday Mode, and takes the order from the footer.

Privacy rule for this script: it never prints or stores question or answer
text. The summary is names and slide numbers only.

Merging caveats (python-pptx has no native slide copy, so this copies the XML):
  * slides land on a blank layout of the merged deck, so a background that
    lived in the guest's theme/master is lost — a picture placed ON the slide
    survives, a theme background does not
  * theme colours resolve against the merged deck's theme, so an "accent 1"
    blue could come out a different blue; explicit RGB colours are untouched
  * text sitting in a title/body placeholder keeps its position and words;
    its font size lived in the layout, so titles are set to 36pt and body
    text to 20pt unless the guest set a size explicitly
  * different slide sizes are scaled to fit and centred; fonts are not scaled
  * animations and transitions are dropped; animated GIFs still play
Open the merged deck in PowerPoint before the party and eyeball it.
"""

import argparse
import copy
import json
import random
import re
import sys
from collections import Counter
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
except ImportError:  # pragma: no cover
    sys.exit("python-pptx is required:  pip install python-pptx")

MARKER = "Reilly Birthday Mode"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SKIP_RELTYPES = ("/slideLayout", "/notesSlide", "/slideMaster", "/theme")

# words people put in filenames that aren't their name
NOISE = re.compile(r"\b(trivia|question|questions|quiz|q|slides?|one|two|three|single|round|pptx|final|v\d+)\b", re.I)


# ---------------------------------------------------------------- names + order
def name_from_file(path):
    """'bob - trivia.pptx' → 'Bob', "Carol's question.pptx" → 'Carol'."""
    s = path.stem
    s = re.sub(r"[_\-–—.()\[\]]+", " ", s)  # separators first, so noise words get word boundaries
    s = re.sub(r"'s\b", "", s)
    s = NOISE.sub(" ", s)
    s = re.sub(r"\d+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return " ".join(w[:1].upper() + w[1:] for w in s.split()) or path.stem


def norm(s):
    return re.sub(r"[^a-z0-9]", "", s.lower())


def decide_order(entries, how, seed, first, last):
    """entries: list of {name, path}. Returns them reordered."""
    by_norm = {norm(e["name"]): e for e in entries}
    if how == "alpha":
        ordered = sorted(entries, key=lambda e: e["name"].lower())
    elif how == "files":
        ordered = sorted(entries, key=lambda e: e["path"].name.lower())
    elif how == "random":
        ordered = list(entries)
        random.Random(seed).shuffle(ordered)
    else:  # explicit comma list; anyone not listed goes after, in file order
        wanted = [norm(x) for x in how.split(",") if x.strip()]
        missing = [x for x in wanted if x not in by_norm]
        if missing:
            sys.exit(f"--order names not found in folder: {', '.join(missing)}")
        listed = [by_norm[x] for x in wanted]
        rest = [e for e in entries if norm(e["name"]) not in wanted]
        ordered = listed + sorted(rest, key=lambda e: e["path"].name.lower())
    for pin, front in ((first, True), (last, False)):
        if not pin:
            continue
        key = norm(pin)
        if key not in by_norm:
            sys.exit(f"--{'first' if front else 'last'} {pin!r} is not one of the decks")
        e = by_norm[key]
        ordered.remove(e)
        ordered.insert(0, e) if front else ordered.append(e)
    return ordered


# ---------------------------------------------------------------- slide copy
def materialise_placeholders(slide):
    """A placeholder with no explicit position inherits it from the layout, which
    we are about to leave behind. Reading .left etc. resolves the inheritance;
    writing it back makes the position explicit on the shape."""
    for shp in slide.shapes:
        if shp.is_placeholder:
            try:
                shp.left, shp.top, shp.width, shp.height = shp.left, shp.top, shp.width, shp.height
            except Exception:
                pass


def ensure_geometry(el):
    """A placeholder inherits its shape geometry from the layout. Once unbound it
    needs its own, or PowerPoint sees a malformed shape. A plain rectangle is
    what title/body placeholders are anyway."""
    spPr = el.find(f"./{{{NS_P}}}spPr")
    if spPr is None:
        return
    if spPr.find(f"{{{NS_A}}}prstGeom") is not None or spPr.find(f"{{{NS_A}}}custGeom") is not None:
        return
    from lxml import etree
    geom = etree.SubElement(spPr, f"{{{NS_A}}}prstGeom", prst="rect")
    etree.SubElement(geom, f"{{{NS_A}}}avLst")
    xfrm = spPr.find(f"{{{NS_A}}}xfrm")
    if xfrm is not None:  # schema order: xfrm, then geometry
        spPr.remove(geom)
        xfrm.addnext(geom)


def default_font_size(el, sz):
    """Set an explicit size (in hundredths of a point) on every run that has
    none, so text that inherited its size from a placeholder doesn't shrink to
    the 18pt default once the placeholder is gone."""
    for tag in ("r", "endParaRPr", "fld"):
        for node in el.iter(f"{{{NS_A}}}{tag}"):
            rpr = node if tag == "endParaRPr" else node.find(f"{{{NS_A}}}rPr")
            if rpr is None:
                from lxml import etree
                rpr = etree.Element(f"{{{NS_A}}}rPr")
                node.insert(0, rpr)
            if rpr.get("sz") is None:
                rpr.set("sz", str(sz))


def remap_rids(el, rid_map):
    for e in el.iter():
        for attr in list(e.attrib):
            if attr.startswith("{" + NS_R + "}") and e.attrib[attr] in rid_map:
                e.attrib[attr] = rid_map[e.attrib[attr]]


def scale_shape(el, factor, dx, dy):
    """Scale a top-level shape's own transform (not group children, whose
    coordinates live in the group's child space)."""
    xfrm = None
    for path in (f"./{{{NS_P}}}spPr/{{{NS_A}}}xfrm", f"./{{{NS_P}}}grpSpPr/{{{NS_A}}}xfrm", f"./{{{NS_P}}}xfrm"):
        xfrm = el.find(path)
        if xfrm is not None:
            break
    if xfrm is None:
        return
    off = xfrm.find(f"{{{NS_A}}}off")
    ext = xfrm.find(f"{{{NS_A}}}ext")
    if off is not None:
        off.set("x", str(int(int(off.get("x")) * factor + dx)))
        off.set("y", str(int(int(off.get("y")) * factor + dy)))
    if ext is not None:
        ext.set("cx", str(int(int(ext.get("cx")) * factor)))
        ext.set("cy", str(int(int(ext.get("cy")) * factor)))


def copy_slide(src_slide, src_prs, dst_prs):
    layout = dst_prs.slide_layouts[6]  # Blank
    dst = dst_prs.slides.add_slide(layout)
    for ph in list(dst.placeholders):
        ph._element.getparent().remove(ph._element)

    # geometry: fit the source slide into the destination, centred
    factor = min(dst_prs.slide_width / src_prs.slide_width, dst_prs.slide_height / src_prs.slide_height)
    dx = int((dst_prs.slide_width - src_prs.slide_width * factor) / 2)
    dy = int((dst_prs.slide_height - src_prs.slide_height * factor) / 2)
    resize = abs(factor - 1) > 1e-6 or dx or dy

    # relationships: pictures, media, hyperlinks. Layout/notes/master stay behind.
    rid_map = {}
    for rel in src_slide.part.rels.values():
        if any(rel.reltype.endswith(t) for t in SKIP_RELTYPES):
            continue
        if rel.is_external:
            new = dst.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            new = dst.part.rels.get_or_add(rel.reltype, rel.target_part)
        rid_map[rel.rId] = new

    # explicit slide background, if the guest set one on the slide itself
    src_cSld = src_slide._element.find(f"{{{NS_P}}}cSld")
    bg = src_cSld.find(f"{{{NS_P}}}bg") if src_cSld is not None else None
    if bg is not None:
        nb = copy.deepcopy(bg)
        remap_rids(nb, rid_map)
        dst._element.find(f"{{{NS_P}}}cSld").insert(0, nb)

    materialise_placeholders(src_slide)
    for shp in src_slide.shapes:
        el = copy.deepcopy(shp._element)
        ph_type = None
        for ph in list(el.iter(f"{{{NS_P}}}ph")):  # unbind from a layout that isn't coming
            ph_type = ph.get("type", "body")
            ph.getparent().remove(ph)
        if ph_type is not None:
            ensure_geometry(el)
            # font size also lived in the layout; give title text a title size
            default_font_size(el, 3600 if ph_type in ("title", "ctrTitle") else 2000)
        remap_rids(el, rid_map)
        if resize:
            scale_shape(el, factor, dx, dy)
        dst.shapes._spTree.insert_element_before(el, "p:extLst")
    return dst


# ---------------------------------------------------------------- index slide
def add_index_slide(prs, title, ordered, seed_note):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    W, H = prs.slide_width, prs.slide_height
    m = int(W * 0.06)

    tb = s.shapes.add_textbox(m, int(H * 0.06), W - 2 * m, int(H * 0.16))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    # the order, in two columns if it's long
    names = [f"{i + 1}.  {e['name']}" for i, e in enumerate(ordered)]
    cols = 2 if len(names) > 8 else 1
    per = -(-len(names) // cols)
    col_w = (W - 2 * m) // cols
    for c in range(cols):
        box = s.shapes.add_textbox(m + c * col_w, int(H * 0.25), col_w, int(H * 0.6))
        tf = box.text_frame
        tf.word_wrap = True
        for i, line in enumerate(names[c * per : (c + 1) * per]):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = line
            para.font.size = Pt(24 if len(names) <= 12 else 18)

    # machine-readable footer the host app parses; small and grey, still honest
    foot = " · ".join(f"{e['name']} {e['slide']}" + (f"–{e['answerSlide']}" if e["answerSlide"] else "") for e in ordered)
    fb = s.shapes.add_textbox(m, int(H * 0.9), W - 2 * m, int(H * 0.08))
    fp = fb.text_frame.paragraphs[0]
    fp.text = f"{MARKER} · {foot}"  # marker box holds ONLY the marker and entries — the app parses every segment
    fp.font.size = Pt(9)
    fp.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    fb.text_frame.word_wrap = True
    if seed_note:  # its own box, so it can never be mistaken for a guest entry
        nb = s.shapes.add_textbox(m, int(H * 0.95), W - 2 * m, int(H * 0.05))
        np_ = nb.text_frame.paragraphs[0]
        np_.text = seed_note
        np_.font.size = Pt(8)
        np_.font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
    return s


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="Merge guests' two-slide trivia decks for Reilly Birthday Mode.")
    ap.add_argument("folder", type=Path, help="folder of one .pptx per guest, named after them")
    ap.add_argument("--out", type=Path, help="output .pptx (default: <folder>/reilly_birthday.pptx)")
    ap.add_argument("--order", default="random", help="random (default) | alpha | files | 'Name,Name,…'")
    ap.add_argument("--seed", type=int, help="seed for --order random; printed either way so the draw is reproducible")
    ap.add_argument("--first", help="pin someone to go first")
    ap.add_argument("--last", help="pin someone to go last (the birthday person, say)")
    ap.add_argument("--points", type=int, default=8, help="points per question (everyone's is worth the same)")
    ap.add_argument("--guesses", type=int, default=2, choices=range(1, 7), help="guesses per player per question, 1–6")
    ap.add_argument("--title", default="Reilly Birthday Trivia")
    a = ap.parse_args()

    if not a.folder.is_dir():
        sys.exit(f"not a folder: {a.folder}")
    out = a.out or a.folder / "reilly_birthday.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    files = sorted(
        p for p in a.folder.glob("*.pptx")
        if not p.name.startswith("~$") and p.resolve() != out.resolve()
    )
    if not files:
        sys.exit("no .pptx files found")

    entries = [{"name": name_from_file(p), "path": p} for p in files]
    dupes = [n for n, c in Counter(norm(e["name"]) for e in entries).items() if c > 1]
    if dupes:
        sys.exit(f"two decks resolve to the same name: {', '.join(dupes)} — rename one")

    seed = a.seed if a.seed is not None else random.randrange(1, 10**6)
    ordered = decide_order(entries, a.order, seed, a.first, a.last)

    # open every deck once; validate; remember question + answer slides
    warnings = []
    for e in ordered:
        prs = Presentation(str(e["path"]))
        n = len(prs.slides)
        if n == 0:
            sys.exit(f"{e['path'].name}: no slides")
        if n > 2:
            warnings.append(f"{e['path'].name}: {n} slides — using the first as the question and the last as the answer, dropping the middle")
        e["prs"] = prs
        e["q"] = prs.slides[0]
        e["a"] = prs.slides[n - 1] if n > 1 else None
        if e["a"] is None:
            warnings.append(f"{e['path'].name}: only one slide, so there is no answer slide")

    # merged deck takes the most common slide size among the inputs
    sizes = Counter((e["prs"].slide_width, e["prs"].slide_height) for e in ordered)
    W, H = sizes.most_common(1)[0][0]
    merged = Presentation()
    merged.slide_width, merged.slide_height = W, H
    if len(sizes) > 1:
        warnings.append(f"{len(sizes)} different slide sizes; using {Emu(W).inches:.2f}×{Emu(H).inches:.2f} in and scaling the rest to fit")

    # slide numbers: index is 1, then pairs
    n = 2
    for e in ordered:
        e["slide"] = n
        e["answerSlide"] = n + 1 if e["a"] is not None else None
        n += 2 if e["a"] is not None else 1

    add_index_slide(merged, a.title, ordered, f"order {a.order}" + (f" seed {seed}" if a.order == "random" else ""))
    for e in ordered:
        copy_slide(e["q"], e["prs"], merged)
        if e["a"] is not None:
            copy_slide(e["a"], e["prs"], merged)
    merged.save(str(out))

    pack = {
        "title": a.title,
        "mode": "birthday",
        "order": a.order,
        "seed": seed if a.order == "random" else None,
        "questions": [
            {
                "id": f"q{i + 1}",
                "author": e["name"],
                "slide": e["slide"],
                "answerSlide": e["answerSlide"],
                "prompt": "",  # deliberately empty: the host is playing too
                "category": "",
                "answers": [],  # deliberately empty: the author grades by hand
                "kind": "single",
                "points": a.points,
                "guesses": a.guesses,
                "autograde": False,
                "threshold": "auto",
                "flags": [],
            }
            for i, e in enumerate(ordered)
        ],
    }
    pack_path = out.with_suffix(".pack.json")
    pack_path.write_text(json.dumps(pack, indent="\t", ensure_ascii=False) + "\n")

    print(f"{len(ordered)} guests → {out.name} ({len(merged.slides)} slides) + {pack_path.name}")
    print(f"order: {a.order}" + (f", seed {seed}  (re-run with --seed {seed} to get the same draw)" if a.order == "random" else ""))
    print(f"{'#':>3}  {'slides':>7}  name  (from file)")
    for i, e in enumerate(ordered, 1):
        sl = f"{e['slide']}–{e['answerSlide']}" if e["answerSlide"] else str(e["slide"])
        print(f"{i:>3}  {sl:>7}  {e['name']}  ({e['path'].name})")
    for w in warnings:
        print("!!", w)


if __name__ == "__main__":
    main()
