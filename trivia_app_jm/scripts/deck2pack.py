#!/usr/bin/env python3
"""deck2pack.py — turn a trivia .pptx into a question pack and slide thumbnails.

    python3 scripts/deck2pack.py "US-UK Trivia.pptx"            # → ./US-UK Trivia.pack/
    python3 scripts/deck2pack.py deck.pptx --out mypack --dpi 36
    python3 scripts/deck2pack.py deck.pptx --no-thumbs          # pack.json only

Writes:
    <out>/pack.json          host-client question pack (load it from the slide tracker)
    <out>/thumbs/slide-NN.png   one low-res PNG per slide, numbered like the deck
    <out>/deck.pdf           intermediate render, kept for eyeballing

Why a local script: the app runs on Cloudflare Workers, which cannot run
LibreOffice, so slide rendering has to happen on the host's own machine (spec:
"likeliest path is host-exported images"). Text extraction is also done here so
the two outputs come from one place, but the host client can do the text part
on its own by unzipping the .pptx in the browser (see mockups/host.html) — the
pairing rules below are mirrored there. Keep them in sync.

Deck conventions this understands (from the US-UK deck):
  * a question is a PAIR of consecutive slides with the same question text; the
    first is shown while people answer, the second reveals the answer
  * the answer is a text box starting with "Answer:" (present on both slides —
    a picture covers it on the question slide)
  * a short extra text box (e.g. "Trivia by the numbers") is a category label
Anything else gets the question with a flag on it rather than being dropped, so
the host can fix or remove it in the slide tracker.
"""

import argparse
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

try:
    from pptx import Presentation
except ImportError:  # pragma: no cover
    sys.exit("python-pptx is required:  pip install python-pptx")

SOFFICE_CANDIDATES = [
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "soffice",
    "libreoffice",
]
ANSWER_RE = re.compile(r"^\s*answer\s*:\s*", re.I)
POINTS_RE = re.compile(r"for\s+(\d+)\s+points?", re.I)


# ---------------------------------------------------------------- extraction
def slide_records(prs):
    """One dict per slide, in deck order: text boxes (paragraphs joined with
    ' / '), picture count, layout name."""
    out = []
    for n, s in enumerate(prs.slides, 1):
        shapes = []
        pics = 0
        for sh in s.shapes:
            if sh.shape_type == 13:  # PICTURE
                pics += 1
            if sh.has_text_frame:
                t = " / ".join(p.text.strip() for p in sh.text_frame.paragraphs if p.text.strip())
                if t:
                    shapes.append(t)
        out.append({"slide": n, "layout": s.slide_layout.name, "shapes": shapes, "pics": pics})
    return out


def norm(s):
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", "", s.lower())).strip()


def split_shapes(shapes):
    """→ (question_text, category_text, answer_text|None). The answer is the
    'Answer:' box; the question is the longest of what's left; anything else is
    a category label."""
    answer = None
    rest = []
    for t in shapes:
        if answer is None and ANSWER_RE.match(t):
            answer = ANSWER_RE.sub("", t).strip()
        else:
            rest.append(t)
    if not rest:
        return "", "", answer
    q = max(rest, key=len)
    cat = " · ".join(t for t in rest if t is not q)
    return q, cat, answer


def same_question(a, b):
    """Pair rule: identical after normalisation, or a long shared prefix (the
    deck had one pair that differed by a single word)."""
    a, b = norm(a), norm(b)
    if not a or not b:
        return False
    if a == b:
        return True
    k = min(len(a), len(b))
    return k >= 30 and a[:30] == b[:30] and (a[:k] == b[:k] or k >= 0.6 * max(len(a), len(b)))


def shared_shapes(a, b):
    """Text boxes slide `a` and slide `b` have in common (ignoring Answer:
    boxes). Two consecutive slides sharing a box are a question pair."""
    return [t for t in a["shapes"] if not ANSWER_RE.match(t) and any(not ANSWER_RE.match(u) and same_question(t, u) for u in b["shapes"])]


def looks_like_question(rec):
    return any("?" in t or "point" in t.lower() for t in rec["shapes"])


def build_pack(slides, title):
    questions = []
    roles = {}  # slide number → question | answer | other
    i = 0
    while i < len(slides):
        s = slides[i]
        nxt = slides[i + 1] if i + 1 < len(slides) else None
        q_text, cat, ans = split_shapes(s["shapes"])
        flags = []
        answer_slide = None

        shared = shared_shapes(s, nxt) if nxt else []
        if shared:
            # A pair. The longest shared box is the question, the other shared
            # boxes are category labels, and if there's no "Answer:" box then
            # whatever the reveal slide has left over is the best guess.
            q_text = max(shared, key=len)
            ans = split_shapes(nxt["shapes"])[2] or ans
            before = [t for t in shared if t is not q_text and s["shapes"].index(t) < s["shapes"].index(q_text)]
            after = [t for t in shared if t is not q_text and s["shapes"].index(t) > s["shapes"].index(q_text)]
            cat = " · ".join(before)  # labels sit above the question
            if ans is None:
                # No "Answer:" box. Text the reveal slide adds is the answer; failing
                # that, the answer box is the last one below the question (it's on
                # both slides — a picture covers it on the question slide).
                extra = [t for t in nxt["shapes"] if not any(same_question(t, u) for u in shared)]
                ans = extra[-1] if extra else (after[-1] if after else "")
                flags.append("no-answer-marker")
            elif after:
                cat = " · ".join(before + after)
            answer_slide = nxt["slide"]
            step = 2
        elif ans is not None or looks_like_question(s):
            flags.append("unpaired")
            if ans is None:
                flags.append("no-answer-marker")
                ans = ""
            step = 1
        else:
            roles[s["slide"]] = "other"
            i += 1
            continue

        answers = [ans] if ans else []
        if " - " in ans:  # "Song - Artist": accept either half too
            answers += [p.strip() for p in ans.split(" - ") if p.strip()]
        if ans.count(",") >= 2:
            flags.append("list-answer")
        if s["pics"] >= 3 or q_text.lower().startswith("picture id"):
            flags.append("picture-id")
        pts = POINTS_RE.findall(q_text + " " + cat)
        if len(pts) > 1:
            flags.append("multi-part")
        points = int(pts[0]) if len(pts) == 1 else 10

        qid = f"q{len(questions) + 1}"
        questions.append(
            {
                "id": qid,
                "slide": s["slide"],
                "answerSlide": answer_slide,
                "prompt": q_text,
                "category": cat,
                "answers": answers,
                "points": points,
                "autograde": True,
                "secondGuess": "grant",
                "threshold": "auto",
                "flags": flags,
            }
        )
        roles[s["slide"]] = "question"
        if answer_slide:
            roles[answer_slide] = "answer"
        i += step

    deck = [
        {"slide": r["slide"], "role": roles.get(r["slide"], "other"), "text": r["shapes"], "thumb": f"thumbs/slide-{r['slide']:02d}.png"}
        for r in slides
    ]
    return {"title": title, "questions": questions, "deck": deck}


# ---------------------------------------------------------------- thumbnails
def find_soffice():
    for c in SOFFICE_CANDIDATES:
        if Path(c).exists() or shutil.which(c):
            return c
    return None


def render_thumbs(pptx, out, dpi):
    soffice = find_soffice()
    if not soffice:
        print("!! LibreOffice not found — skipped thumbnails. Install it, or export slides as images from PowerPoint.", file=sys.stderr)
        return 0
    if not shutil.which("pdftoppm"):
        print("!! pdftoppm not found (brew install poppler) — skipped thumbnails.", file=sys.stderr)
        return 0
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out), str(pptx)], check=True, capture_output=True)
    pdf = out / (pptx.stem + ".pdf")
    pdf.rename(out / "deck.pdf")
    thumbs = out / "thumbs"
    thumbs.mkdir(exist_ok=True)
    for old in thumbs.glob("slide-*.png"):
        old.unlink()
    subprocess.run(["pdftoppm", "-r", str(dpi), "-png", str(out / "deck.pdf"), str(thumbs / "slide")], check=True)
    return len(list(thumbs.glob("slide-*.png")))


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("pptx", type=Path)
    ap.add_argument("--out", type=Path, help="output folder (default: <deck name>.pack next to the deck)")
    ap.add_argument("--dpi", type=int, default=24, help="thumbnail resolution; 24 ≈ 240px wide for a 16:9 deck")
    ap.add_argument("--no-thumbs", action="store_true")
    a = ap.parse_args()

    if not a.pptx.exists():
        sys.exit(f"no such file: {a.pptx}")
    out = a.out or a.pptx.with_suffix(".pack")
    out.mkdir(parents=True, exist_ok=True)

    prs = Presentation(str(a.pptx))
    slides = slide_records(prs)
    pack = build_pack(slides, a.pptx.stem)
    (out / "pack.json").write_text(json.dumps(pack, indent="\t", ensure_ascii=False) + "\n")

    n_thumbs = 0 if a.no_thumbs else render_thumbs(a.pptx, out, a.dpi)

    print(f"{len(slides)} slides → {len(pack['questions'])} questions, {n_thumbs} thumbnails → {out}/")
    print(f"{'#':>3} {'slide':>5}  {'pts':>3}  flags / prompt")
    for i, q in enumerate(pack["questions"], 1):
        sl = f"{q['slide']}→{q['answerSlide']}" if q["answerSlide"] else str(q["slide"])
        flag = ("⚠ " + ", ".join(q["flags"]) + "  ") if q["flags"] else ""
        print(f"{i:>3} {sl:>5}  {q['points']:>3}  {flag}{q['prompt'][:70]}")
        print(f"{'':>15}  = {' | '.join(q['answers'])[:90]}")
    skipped = [d["slide"] for d in pack["deck"] if d["role"] == "other"]
    if skipped:
        print(f"not questions: slides {', '.join(map(str, skipped))}")


if __name__ == "__main__":
    main()
